#!/usr/bin/env python3
"""
Times New Roman Font PPT
Heading: 24pt, Body: 20pt
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

BLACK = RGBColor(0, 0, 0)
DARK_GRAY = RGBColor(50, 50, 50)

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_simple_slide(prs, title, content_lines):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title - 24pt
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    
    # Simple line
    line = slide.shapes.add_shape(1, Inches(0.5), Inches(1.2), Inches(2), Pt(1))
    line.fill.solid()
    line.fill.fore_color.rgb = DARK_GRAY
    line.line.fill.background()
    
    # Content - 20pt
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.333), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    for i, line_text in enumerate(content_lines):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        
        p.text = line_text
        p.font.size = Pt(20)
        p.font.color.rgb = BLACK
        p.font.name = "Times New Roman"
        p.space_after = Pt(6)
    
    return slide

def add_title_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.2))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "Job Tracking System"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Full Stack Web Application"
    p2.font.size = Pt(20)
    p2.font.color.rgb = DARK_GRAY
    p2.font.name = "Times New Roman"
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(12)
    
    p3 = tf.add_paragraph()
    p3.text = "React | Node.js | Express | MongoDB Atlas"
    p3.font.size = Pt(16)
    p3.font.color.rgb = DARK_GRAY
    p3.font.name = "Times New Roman"
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(20)
    
    p4 = tf.add_paragraph()
    p4.text = "Presented by: [Your Name] | April 2026"
    p4.font.size = Pt(12)
    p4.font.color.rgb = DARK_GRAY
    p4.font.name = "Times New Roman"
    p4.alignment = PP_ALIGN.CENTER
    p4.space_before = Pt(30)
    
    return slide

def add_two_column_slide(prs, title, left_title, left_items, right_title, right_items):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title - 24pt
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    
    # Left title - 22pt
    lt = slide.shapes.add_textbox(Inches(0.5), Inches(1.4), Inches(6), Inches(0.5))
    tf = lt.text_frame
    p = tf.paragraphs[0]
    p.text = left_title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    
    # Left content - 20pt
    lc = slide.shapes.add_textbox(Inches(0.5), Inches(2), Inches(6), Inches(5))
    tf = lc.text_frame
    tf.word_wrap = True
    
    for item in left_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.font.color.rgb = BLACK
        p.font.name = "Times New Roman"
        p.space_after = Pt(4)
    
    # Right title - 22pt
    rt = slide.shapes.add_textbox(Inches(6.8), Inches(1.4), Inches(6), Inches(0.5))
    tf = rt.text_frame
    p = tf.paragraphs[0]
    p.text = right_title
    p.font.size = Pt(22)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    
    # Right content - 20pt
    rc = slide.shapes.add_textbox(Inches(6.8), Inches(2), Inches(6), Inches(5))
    tf = rc.text_frame
    tf.word_wrap = True
    
    for item in right_items:
        p = tf.add_paragraph()
        p.text = item
        p.font.size = Pt(20)
        p.font.color.rgb = BLACK
        p.font.name = "Times New Roman"
        p.space_after = Pt(4)
    
    return slide

def add_table_slide(prs, title, headers, rows):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    # Title - 24pt
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.4), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(24)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    
    # Table
    num_rows = len(rows) + 1
    num_cols = len(headers)
    
    table = slide.shapes.add_table(num_rows, num_cols, Inches(0.5), Inches(1.4), Inches(12.333), Inches(0.55 * num_rows)).table
    
    # Headers - 18pt
    for i, header_text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header_text
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.bold = True
            paragraph.font.size = Pt(18)
            paragraph.font.name = "Times New Roman"
    
    # Rows - 16pt
    for row_idx, row_data in enumerate(rows, 1):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_text)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(16)
                paragraph.font.name = "Times New Roman"
    
    return slide

def add_end_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    
    title = slide.shapes.add_textbox(Inches(0.5), Inches(2.8), Inches(12.333), Inches(1))
    tf = title.text_frame
    p = tf.paragraphs[0]
    p.text = "THANK YOU"
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = BLACK
    p.font.name = "Times New Roman"
    p.alignment = PP_ALIGN.CENTER
    
    p2 = tf.add_paragraph()
    p2.text = "Questions?"
    p2.font.size = Pt(24)
    p2.font.color.rgb = DARK_GRAY
    p2.font.name = "Times New Roman"
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)
    
    p3 = tf.add_paragraph()
    p3.text = "GitHub: [Repository URL]"
    p3.font.size = Pt(16)
    p3.font.color.rgb = DARK_GRAY
    p3.font.name = "Times New Roman"
    p3.alignment = PP_ALIGN.CENTER
    p3.space_before = Pt(30)
    
    return slide

def main():
    prs = create_presentation()
    
    # 14 slides
    add_title_slide(prs)
    
    add_two_column_slide(prs, "The Challenge - Problem Statement",
        "For Job Seekers:",
        ["- Managing 50+ applications across multiple companies",
         "- Missing deadlines and follow-up dates",
         "- Scattered notes and spreadsheets",
         "- No visibility into application progress"],
        "For Recruiters:",
        ["- Overwhelmed by application volume",
         "- Disorganized candidate tracking",
         "- No analytics on hiring pipeline",
         "- Manual document management"])
    
    add_simple_slide(prs, "Our Solution: Job Tracking System", [
        "A unified platform connecting job seekers, recruiters, and admins",
        "",
        "Key Capabilities:",
        "- Track job applications with status updates",
        "- Build professional resumes with structured sections",
        "- Apply to jobs through integrated portal",
        "- Analyze hiring metrics with visual dashboards",
        "- Extract skills from job descriptions",
        "",
        "User Roles:",
        "- Job Seeker: Track apps, build resumes, apply for jobs",
        "- Recruiter: Post jobs, review candidates, analytics",
        "- Admin: Manage users, oversee platform"
    ])
    
    add_simple_slide(prs, "Technologies Used", [
        "Frontend:",
        "- React 19 - Component-based UI library",
        "- React Router - Client-side navigation",
        "- Chart.js & Recharts - Data visualization",
        "- Axios - HTTP client for API calls",
        "",
        "Backend:",
        "- Node.js - JavaScript runtime",
        "- Express.js - Web framework for REST APIs",
        "- MongoDB Driver - Database connectivity",
        "",
        "Database:",
        "- MongoDB Atlas - Cloud-hosted NoSQL database",
        "- 6 collections with optimized indexing"
    ])
    
    add_simple_slide(prs, "3-Tier System Architecture", [
        "PRESENTATION LAYER:",
        "- React Frontend (Port 3000)",
        "  - Dashboard Components",
        "  - Resume Builder Forms",
        "  - Job Portal Interface",
        "",
        "APPLICATION LAYER:",
        "- Express Server (Port 5000)",
        "  - 30+ REST API Endpoints",
        "  - Middleware (CORS, Auth)",
        "  - Business Logic Handlers",
        "",
        "DATA LAYER:",
        "- MongoDB Atlas (Cloud)",
        "  - users, jobs, resumes",
        "  - job_postings, applications"
    ])
    
    add_two_column_slide(prs, "Core Modules: Authentication & Job Tracking",
        "Authentication Module:",
        ["- User registration with role selection",
         "- Login with credential validation",
         "- Session management via localStorage",
         "- Role-based access control",
         "",
         "Job Tracking Module:",
         "- POST /jobs -> Add application",
         "- GET /jobs -> View all applications",
         "- PUT /jobs/:id -> Update status",
         "- DELETE /jobs/:id -> Remove application"],
        "Fields Tracked:",
        ["- Company name",
         "- Role / Position",
         "- Application Status",
         "- Deadline",
         "- Follow-up date",
         "- Notes",
         "- Job URL"])
    
    add_simple_slide(prs, "Core Modules: Resume Builder & Job Portal", [
        "Resume Builder - 6 Sections:",
        "1. Personal Information",
        "2. Professional Summary",
        "3. Work Experience (multiple entries)",
        "4. Education History",
        "5. Skills (tag-based)",
        "6. Projects Portfolio",
        "",
        "Job Portal - For Recruiters:",
        "- Create job postings with requirements",
        "- Review applicants with documents",
        "- Update application status (accept/reject)",
        "- View hiring analytics",
        "",
        "Job Portal - For Job Seekers:",
        "- Browse available positions",
        "- Apply with resume + cover letter",
        "- Track application status",
        "- Withdraw applications"
    ])
    
    add_simple_slide(prs, "Core Modules: Admin, File Upload & Skill Analysis", [
        "Admin Module:",
        "- View all users and job postings",
        "- Delete users (cascades to their data)",
        "- Manage platform content",
        "- View system-wide statistics",
        "",
        "File Upload Module:",
        "- Base64 encoding for storage",
        "- Resume and cover letter attachments",
        "- 10MB size limit",
        "- Retrieval for recruiter review",
        "",
        "Skill Analysis Module:",
        "- Keyword-based skill extraction",
        "- 40+ predefined skills database",
        "- Contextual analysis (remote -> Remote Collaboration)",
        "- Gap identification for job seekers"
    ])
    
    add_simple_slide(prs, "Major API Endpoints", [
        "Authentication:",
        "- POST /signup - User registration",
        "- POST /login - User authentication",
        "",
        "Job Management:",
        "- POST /jobs - Add application",
        "- GET /jobs - List applications",
        "- PUT /jobs/:id - Update application",
        "- DELETE /jobs/:id - Delete application",
        "",
        "Job Portal:",
        "- POST /create-job-posting - Recruiter posts job",
        "- POST /apply-job - Seeker applies",
        "- PATCH /job-applications/:id/status - Recruiter decision"
    ])
    
    add_table_slide(prs, "MongoDB Collections",
        ["Collection", "Purpose", "Key Fields"],
        [
            ["users", "Authentication", "username, password, role"],
            ["jobs", "Personal job apps", "user_id, company, role, status, deadline"],
            ["resumes", "Resume builder data", "user_id, personalInfo, experience[], skills[]"],
            ["job_postings", "Recruiter listings", "recruiter_id, position, company, requirements[]"],
            ["job_applications", "Portal applications", "user_id, job_id, resume_attachment, status"],
            ["job_seeker_resumes", "Uploaded files", "user_id, name, file_content (base64)"]
        ]
    )
    
    add_simple_slide(prs, "Key Features", [
        "1. Visual Dashboard:",
        "   - Charts showing application status distribution",
        "   - Deadline calendar with color coding",
        "   - Statistics cards (total, interviews, offers)",
        "",
        "2. Resume Builder:",
        "   - Multiple resume versions",
        "   - Dynamic section addition",
        "   - Professional structure",
        "",
        "3. Job Portal:",
        "   - Direct seeker-recruiter connection",
        "   - Document upload and review",
        "   - Real-time status updates",
        "",
        "4. Analytics:",
        "   - Recruiter hiring metrics",
        "   - Application conversion rates",
        "   - Pipeline visualization"
    ])
    
    add_simple_slide(prs, "Demo Walkthrough", [
        "Job Seeker Journey:",
        "1. Sign up -> Create account as job seeker",
        "2. Build resume -> Fill 6 sections with experience",
        "3. Add job -> Track Google application with deadline",
        "4. Browse portal -> Find matching positions",
        "5. Apply -> Upload resume + cover letter",
        "6. Track progress -> View status in dashboard",
        "",
        "Recruiter Journey:",
        "1. Post job -> Create Full Stack Developer listing",
        "2. Review applicants -> View 20+ applications",
        "3. Access documents -> Download resumes",
        "4. Make decisions -> Accept/reject candidates",
        "5. View analytics -> Check conversion rates"
    ])
    
    add_simple_slide(prs, "Summary & Next Steps", [
        "Achievements:",
        "- Full-stack implementation with modern tech",
        "- Multi-role architecture serving 3 user types",
        "- 30+ REST APIs with proper validation",
        "- 6 MongoDB collections with indexing",
        "- File handling, analytics, and skill extraction",
        "",
        "Limitations:",
        "- Passwords stored as plaintext (needs hashing)",
        "- No email notifications",
        "- No real-time updates (WebSocket)",
        "- Base64 storage (not optimal for scale)",
        "",
        "Future Enhancements:",
        "- JWT authentication & password hashing",
        "- Email/SMS notifications",
        "- AI-powered job matching",
        "- LinkedIn integration",
        "- Mobile app (React Native)",
        "- Resume PDF export"
    ])
    
    add_end_slide(prs)
    
    output = r"c:\Users\arebu\Desktop\job-tracker-project\Job_Tracking_System_TimesRoman.pptx"
    prs.save(output)
    print(f"✅ Times New Roman PPT saved: {output}")
    print(f"📊 14 slides - Heading: 24pt, Body: 20pt, Font: Times New Roman")

if __name__ == "__main__":
    main()
