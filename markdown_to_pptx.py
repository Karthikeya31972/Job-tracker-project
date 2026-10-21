#!/usr/bin/env python3
"""
Convert markdown presentation slides to PowerPoint (.pptx) format
"""

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
import re

def create_presentation():
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    return prs

def add_title_slide(prs, title, subtitle=""):
    slide_layout = prs.slide_layouts[6]  # Blank layout
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(25, 55, 109)
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(28)
        p2.font.color.rgb = RGBColor(200, 200, 200)
        p2.alignment = PP_ALIGN.CENTER
    
    return slide

def add_content_slide(prs, title, content_items, subtitle=""):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header bar
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(25, 55, 109)
    header.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    if subtitle:
        p2 = tf.add_paragraph()
        p2.text = subtitle
        p2.font.size = Pt(18)
        p2.font.color.rgb = RGBColor(200, 220, 255)
    
    # Content
    content_box = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(12.333), Inches(5.5))
    tf = content_box.text_frame
    tf.word_wrap = True
    
    first = True
    for item in content_items:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        
        # Check if it's a heading
        if item.startswith('### '):
            p.text = item[4:]
            p.font.size = Pt(24)
            p.font.bold = True
            p.font.color.rgb = RGBColor(25, 55, 109)
            p.space_before = Pt(12)
        elif item.startswith('**') and item.endswith('**'):
            p.text = item.strip('*')
            p.font.size = Pt(20)
            p.font.bold = True
            p.font.color.rgb = RGBColor(40, 80, 140)
        elif item.startswith('> '):
            p.text = '"' + item[2:].strip('"') + '"'
            p.font.size = Pt(18)
            p.font.italic = True
            p.font.color.rgb = RGBColor(80, 80, 80)
            p.space_before = Pt(12)
        elif item.startswith('- '):
            p.text = "• " + item[2:]
            p.font.size = Pt(18)
            p.level = 0
            p.font.color.rgb = RGBColor(50, 50, 50)
        elif item.startswith('  - '):
            p.text = "   ◦ " + item[4:]
            p.font.size = Pt(16)
            p.level = 1
            p.font.color.rgb = RGBColor(70, 70, 70)
        elif item.startswith('1. ') or item.startswith('2. ') or item.startswith('3. ') or item.startswith('4. ') or item.startswith('5. ') or item.startswith('6. '):
            p.text = item
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(50, 50, 50)
        else:
            p.text = item
            p.font.size = Pt(18)
            p.font.color.rgb = RGBColor(50, 50, 50)
        
        p.space_after = Pt(6)
    
    return slide

def add_table_slide(prs, title, headers, rows):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header bar
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(25, 55, 109)
    header.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Table
    num_rows = len(rows) + 1
    num_cols = len(headers)
    
    table = slide.shapes.add_table(
        num_rows, num_cols, Inches(0.5), Inches(1.5), Inches(12.333), Inches(0.6 * num_rows)
    ).table
    
    # Set header
    for i, header_text in enumerate(headers):
        cell = table.cell(0, i)
        cell.text = header_text
        cell.fill.solid()
        cell.fill.fore_color.rgb = RGBColor(40, 80, 140)
        for paragraph in cell.text_frame.paragraphs:
            paragraph.font.color.rgb = RGBColor(255, 255, 255)
            paragraph.font.bold = True
            paragraph.font.size = Pt(14)
    
    # Set rows
    for row_idx, row_data in enumerate(rows, 1):
        for col_idx, cell_text in enumerate(row_data):
            cell = table.cell(row_idx, col_idx)
            cell.text = str(cell_text)
            for paragraph in cell.text_frame.paragraphs:
                paragraph.font.size = Pt(12)
                if row_idx % 2 == 0:
                    cell.fill.solid()
                    cell.fill.fore_color.rgb = RGBColor(240, 245, 250)
    
    return slide

def add_code_slide(prs, title, code_lines):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Header bar
    header = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, Inches(1.2)
    )
    header.fill.solid()
    header.fill.fore_color.rgb = RGBColor(25, 55, 109)
    header.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(0.25), Inches(12.333), Inches(0.8))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(36)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    
    # Code box
    code_box = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, Inches(0.5), Inches(1.5), Inches(12.333), Inches(5.5)
    )
    code_box.fill.solid()
    code_box.fill.fore_color.rgb = RGBColor(30, 30, 30)
    code_box.line.color.rgb = RGBColor(60, 60, 60)
    
    # Code text
    text_box = slide.shapes.add_textbox(Inches(0.7), Inches(1.7), Inches(11.933), Inches(5.1))
    tf = text_box.text_frame
    tf.word_wrap = True
    
    first = True
    for line in code_lines:
        if first:
            p = tf.paragraphs[0]
            first = False
        else:
            p = tf.add_paragraph()
        
        p.text = line
        p.font.size = Pt(14)
        p.font.name = "Consolas"
        p.font.color.rgb = RGBColor(200, 200, 200)
    
    return slide

def add_end_slide(prs, title="THANK YOU", subtitle="Questions?"):
    slide_layout = prs.slide_layouts[6]
    slide = prs.slides.add_slide(slide_layout)
    
    # Background
    background = slide.shapes.add_shape(
        MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height
    )
    background.fill.solid()
    background.fill.fore_color.rgb = RGBColor(25, 55, 109)
    background.line.fill.background()
    
    # Title
    title_box = slide.shapes.add_textbox(Inches(0.5), Inches(2.5), Inches(12.333), Inches(1.5))
    tf = title_box.text_frame
    p = tf.paragraphs[0]
    p.text = title
    p.font.size = Pt(72)
    p.font.bold = True
    p.font.color.rgb = RGBColor(255, 255, 255)
    p.alignment = PP_ALIGN.CENTER
    
    # Subtitle
    p2 = tf.add_paragraph()
    p2.text = subtitle
    p2.font.size = Pt(36)
    p2.font.color.rgb = RGBColor(200, 220, 255)
    p2.alignment = PP_ALIGN.CENTER
    p2.space_before = Pt(20)
    
    return slide

def main():
    prs = create_presentation()
    
    # Slide 1: Title
    add_title_slide(prs, "Job Tracking System", "Full Stack Web Application\n\nTech Stack: React | Node.js | Express | MongoDB Atlas")
    
    # Slide 2: Problem Statement
    add_content_slide(prs, "The Challenge - Problem Statement", [
        "### For Job Seekers:",
        "- Managing 50+ applications across multiple companies",
        "- Missing deadlines and follow-up dates",
        "- Scattered notes and spreadsheets",
        "- No visibility into application progress",
        "",
        "### For Recruiters:",
        "- Overwhelmed by application volume",
        "- Disorganized candidate tracking",
        "- No analytics on hiring pipeline",
        "- Manual document management",
        "",
        "> The average job seeker applies to 50+ jobs but loses track within a week"
    ])
    
    # Slide 3: Solution Overview
    add_content_slide(prs, "Our Solution: Job Tracking System", [
        "**A unified platform connecting job seekers, recruiters, and admins**",
        "",
        "### Key Capabilities:",
        "- **Track** job applications with status updates",
        "- **Build** professional resumes with structured sections",
        "- **Apply** to jobs through integrated portal",
        "- **Analyze** hiring metrics with visual dashboards",
        "- **Extract** skills from job descriptions",
        "",
        "### User Roles:",
        "- **Job Seeker** → Track apps, build resumes, apply for jobs",
        "- **Recruiter** → Post jobs, review candidates, analytics",
        "- **Admin** → Manage users, oversee platform"
    ])
    
    # Slide 4: Tech Stack
    add_content_slide(prs, "Technologies Used", [
        "### Frontend",
        "- **React 19** - Component-based UI library",
        "- **React Router** - Client-side navigation",
        "- **Chart.js & Recharts** - Data visualization",
        "- **Axios** - HTTP client for API calls",
        "",
        "### Backend",
        "- **Node.js** - JavaScript runtime",
        "- **Express.js** - Web framework for REST APIs",
        "- **MongoDB Driver** - Database connectivity",
        "",
        "### Database",
        "- **MongoDB Atlas** - Cloud-hosted NoSQL database",
        "- 6 collections with optimized indexing"
    ])
    
    # Slide 5: System Architecture (as text)
    add_content_slide(prs, "3-Tier System Architecture", [
        "### PRESENTATION LAYER",
        "- React Frontend (Port 3000)",
        "  - Dashboard Components",
        "  - Resume Builder Forms",
        "  - Job Portal Interface",
        "",
        "### APPLICATION LAYER",
        "- Express Server (Port 5000)",
        "  - 30+ REST API Endpoints",
        "  - Middleware (CORS, Auth)",
        "  - Business Logic Handlers",
        "",
        "### DATA LAYER",
        "- MongoDB Atlas (Cloud)",
        "  - users, jobs, resumes",
        "  - job_postings, applications"
    ])
    
    # Slide 6: Core Modules Part 1
    add_content_slide(prs, "Core Modules: Authentication & Job Tracking", [
        "### Authentication Module",
        "- User registration with role selection",
        "- Login with credential validation",
        "- Session management via localStorage",
        "- Role-based access control (client/recruiter/admin)",
        "",
        "### Job Tracking Module",
        "- **POST** /jobs → Add application",
        "- **GET** /jobs → View all applications",
        "- **PUT** /jobs/:id → Update status",
        "- **DELETE** /jobs/:id → Remove application",
        "",
        "**Fields Tracked:** Company, Role, Status, Deadline, Follow-up, Notes, URL"
    ])
    
    # Slide 7: Core Modules Part 2
    add_content_slide(prs, "Core Modules: Resume Builder & Job Portal", [
        "### Resume Builder Module",
        "1. Personal Information",
        "2. Professional Summary",
        "3. Work Experience (multiple entries)",
        "4. Education History",
        "5. Skills (tag-based)",
        "6. Projects Portfolio",
        "",
        "### Job Portal Module",
        "**For Recruiters:**",
        "- Create job postings with requirements",
        "- Review applicants with documents",
        "- Update application status (accept/reject)",
        "- View hiring analytics",
        "",
        "**For Job Seekers:**",
        "- Browse available positions",
        "- Apply with resume + cover letter",
        "- Track application status"
    ])
    
    # Slide 8: Core Modules Part 3
    add_content_slide(prs, "Core Modules: Admin, File Upload & Skill Analysis", [
        "### Admin Module",
        "- View all users and job postings",
        "- Delete users (cascades to their data)",
        "- Manage platform content",
        "- View system-wide statistics",
        "",
        "### File Upload Module",
        "- Base64 encoding for storage",
        "- Resume and cover letter attachments",
        "- 10MB size limit",
        "- Retrieval for recruiter review",
        "",
        "### Skill Analysis Module",
        "- Keyword-based skill extraction",
        "- 40+ predefined skills database",
        "- Contextual analysis (remote → Remote Collaboration)",
        "- Gap identification for job seekers"
    ])
    
    # Slide 9: API Design - Tables
    add_content_slide(prs, "Major API Endpoints", [
        "### Authentication",
        "- **POST** /signup - User registration",
        "- **POST** /login - User authentication",
        "",
        "### Job Management",
        "- **POST** /jobs - Add application",
        "- **GET** /jobs - List applications",
        "- **PUT** /jobs/:id - Update application",
        "- **DELETE** /jobs/:id - Delete application",
        "",
        "### Job Portal",
        "- **POST** /create-job-posting - Recruiter posts job",
        "- **POST** /apply-job - Seeker applies",
        "- **PATCH** /job-applications/:id/status - Recruiter decision"
    ])
    
    # Slide 10: Database Design
    add_table_slide(prs, "MongoDB Collections", 
        ["Collection", "Purpose", "Key Fields"],
        [
            ["users", "Authentication", "username, password, role"],
            ["jobs", "Personal job apps", "user_id, company, role, status, deadline"],
            ["resumes", "Resume builder data", "user_id, personalInfo, experience[], skills[]"],
            ["job_postings", "Recruiter listings", "recruiter_id, position, company, requirements[]"],
            ["job_applications", "Portal applications", "user_id, job_id, resume_attachment, status"],
            ["job_seeker_resumes", "Uploaded files", "user_id, name, file_content (base64)"]
        ]
    )
    
    # Slide 11: Key Features
    add_content_slide(prs, "What Makes This System Powerful", [
        "### 1. Visual Dashboard",
        "- Charts showing application status distribution",
        "- Deadline calendar with color coding",
        "- Statistics cards (total, interviews, offers)",
        "",
        "### 2. Resume Builder",
        "- Multiple resume versions",
        "- Dynamic section addition",
        "- Professional structure",
        "",
        "### 3. Job Portal",
        "- Direct seeker-recruiter connection",
        "- Document upload and review",
        "- Real-time status updates",
        "",
        "### 4. Analytics & Smart Features",
        "- Recruiter hiring metrics",
        "- Skill extraction from descriptions",
        "- Deadline reminders"
    ])
    
    # Slide 12: Demo Walkthrough
    add_content_slide(prs, "System in Action - Demo Walkthrough", [
        "### Job Seeker Journey:",
        "1. **Sign up** → Create account as job seeker",
        "2. **Build resume** → Fill 6 sections with experience",
        "3. **Add job** → Track Google application with deadline",
        "4. **Browse portal** → Find matching positions",
        "5. **Apply** → Upload resume + cover letter",
        "6. **Track progress** → View status in dashboard",
        "",
        "### Recruiter Journey:",
        "1. **Post job** → Create Full Stack Developer listing",
        "2. **Review applicants** → View 20+ applications",
        "3. **Access documents** → Download resumes",
        "4. **Make decisions** → Accept/reject candidates",
        "5. **View analytics** → Check conversion rates"
    ])
    
    # Slide 13: Conclusion
    add_content_slide(prs, "Summary & Next Steps", [
        "### Achievements",
        "- Full-stack implementation with modern tech",
        "- Multi-role architecture serving 3 user types",
        "- 30+ REST APIs with proper validation",
        "- 6 MongoDB collections with indexing",
        "- File handling, analytics, and skill extraction",
        "",
        "### Limitations",
        "- Passwords stored as plaintext (needs hashing)",
        "- No email notifications",
        "- No real-time updates (WebSocket)",
        "",
        "### Future Enhancements",
        "- JWT authentication & password hashing",
        "- Email/SMS notifications",
        "- AI-powered job matching",
        "- LinkedIn integration",
        "- Mobile app (React Native)",
        "- Resume PDF export"
    ])
    
    # Final Slide
    add_end_slide(prs, "THANK YOU", "Questions?\n\nGitHub: [Repository URL]\nDemo: http://localhost:3000")
    
    # Save
    output_path = r"c:\Users\arebu\Desktop\job-tracker-project\Job_Tracking_System_Presentation.pptx"
    prs.save(output_path)
    print(f"Presentation saved to: {output_path}")

if __name__ == "__main__":
    main()
